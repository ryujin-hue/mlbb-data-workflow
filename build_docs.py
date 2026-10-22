from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches as PInches, Pt as PPt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor as PPTColor
from pathlib import Path
import csv, shutil
ROOT=Path(__file__).resolve().parent
REPORT=ROOT/'reports'/'MLBB_Meta_Counter_Full_Extraction_V3.docx'
PRES=ROOT/'reports'/'MLBB_Meta_Counter_Full_Extraction_V3_Presentation.pptx'

NAVY='102338'; BLUE='126E82'; CYAN='38C6D9'; LIGHT='F0F5F8'; DARK='1B2935'; MUTED='566775'; WHITE='FFFFFF'; ORANGE='F2A65A'

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def set_cell_text(cell, text, bold=False, color=DARK, size=9):
    cell.text=''
    p=cell.paragraphs[0]
    p.paragraph_format.space_after=Pt(0)
    r=p.add_run(str(text)); r.bold=bold; r.font.name='Aptos'; r.font.size=Pt(size); r.font.color.rgb=RGBColor.from_string(color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_table(doc, headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i],h,True,WHITE,8.5); shade(t.rows[0].cells[i],NAVY)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,x in enumerate(row):
            set_cell_text(cells[i],x,False,DARK,8)
            if ri%2==1: shade(cells[i],LIGHT)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def set_repeat_table_header(row):
    trPr=row._tr.get_or_add_trPr(); tblHeader=OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'),'true'); trPr.append(tblHeader)

def make_report():
    d=Document(); sec=d.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
    styles=d.styles
    styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(9.5); styles['Normal'].font.color.rgb=RGBColor.from_string(DARK)
    styles['Normal'].paragraph_format.space_after=Pt(5)
    for style, size, col in [('Heading 1',18,BLUE),('Heading 2',13,NAVY),('Heading 3',10.5,BLUE)]:
        s=styles[style]; s.font.name='Aptos Display'; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(col)
    # header/footer
    header=sec.header.paragraphs[0]; header.text='MLBB META & COUNTER INTELLIGENCE  /  V3 EXTRACTION'; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    for r in header.runs: r.font.name='Aptos'; r.font.size=Pt(7); r.font.color.rgb=RGBColor.from_string(MUTED)
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    rr=footer.add_run('Research build • 22 October 2026 • Validate source coverage before production use'); rr.font.size=Pt(7); rr.font.color.rgb=RGBColor.from_string(MUTED)
    # cover
    p=d.add_paragraph(); p.paragraph_format.space_before=Pt(48); p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    r=p.add_run('MLBB'); r.bold=True; r.font.name='Aptos Display'; r.font.size=Pt(42); r.font.color.rgb=RGBColor.from_string(BLUE)
    p=d.add_paragraph(); p.paragraph_format.space_after=Pt(1)
    r=p.add_run('META & COUNTER'); r.bold=True; r.font.name='Aptos Display'; r.font.size=Pt(27); r.font.color.rgb=RGBColor.from_string(NAVY)
    p=d.add_paragraph(); p.paragraph_format.space_after=Pt(20)
    r=p.add_run('FULL PUBLIC-DATA EXTRACTION • V3'); r.bold=True; r.font.size=Pt(14); r.font.color.rgb=RGBColor.from_string(BLUE)
    p=d.add_paragraph('Pipeline, sourced snapshots, data schema, and audit controls for MPL, MSC, M-Series and other professional MLBB tournaments.'); p.paragraph_format.space_after=Pt(18)
    t=d.add_table(rows=1,cols=1); shade(t.cell(0,0),NAVY); set_cell_text(t.cell(0,0),'INTEGRITY STATUS: The attached ZIP is an extraction toolkit with sourced starter snapshots. The live historical extraction has not been run in this environment because direct external network access is unavailable. The collector is designed to run in a network-enabled environment and checks for a real 100 MiB output without filler data.',False,WHITE,11)
    d.add_paragraph()
    d.add_heading('Executive overview',1)
    d.add_paragraph('This V3 package expands the earlier research bundle into a reproducible collection workflow. It preserves already captured, source-labelled hero snapshots and adds an extractor for public tournament pages, table parsing, public repository assets, and run-level audit logs. It does not represent the completed exhaustive archive of every historical tournament or every draft.')
    add_table(d,['Requirement','V3 treatment'],[
        ('All hero coverage','Global pro snapshot included; full hero/season coverage requires validation from source tables.'),
        ('All MPL seasons','Season-by-season inventory and game-level schema provided; historical completeness is not asserted.'),
        ('M-Series / MSC / other events','Recent index counts and MSC 2026 hero snapshot included; extractor discovers accessible event pages.'),
        ('Draft counters','Game-level data model supports heroes, bans, order, result, duration and patch where sourced.'),
        ('100 MiB target','Runtime checks actual output size. If sources yield less than 100 MiB, it reports failure without padding or synthetic data.')
    ],[1.6,5.4])
    d.add_page_break()
    d.add_heading('1. Source inventory and current visible coverage',1)
    d.add_paragraph('The MLBBDex esports index checked on 22 October 2026 lists the current MPL S18 pages plus MSC 2026 and M7. Its page reports 554 recent games across MSC 2026, MPL ID S18, MPL PH S18 and MPL MY S18. These displayed counts are source-index counts, not an independent confirmation that every game has a complete draft record.')
    rows=[]
    with open(ROOT/'data'/'tournament_reference_index.csv',encoding='utf-8-sig',newline='') as f:
        for row in csv.DictReader(f): rows.append((row['tournament'],row['games_recorded'],row['status_as_of_2026-10-22']))
    add_table(d,['Tournament','Games listed','Index status'],rows,[3.0,1.2,2.6])
    d.add_heading('Main sources',2)
    add_table(d,['Source','What it contributes','Important limit'],[
        ('MLBBDex esports index','Tournament discovery, page links, displayed game counts, hero data by event.','Coverage depends on the pages currently indexed.'),
        ('MLBB Analytics','Global professional hero statistics; prior capture reports 2,533 pro games across 20 championships.','Aggregate sample should not be read as every match of every MPL season.'),
        ('MLBBHub MPL MY S18','Visible season-level hero pick/ban and WR rows.','Current snapshot subset, not complete all-season coverage.'),
        ('Liquipedia statistics portal','Broad historical index; its public page reports years 2017–2026, 1,128 tournaments and 72,512 games in its indexed set.','Access may be blocked; collector logs that instead of bypassing it.'),
        ('Public GitHub datasets','Open game reference JSON and images; community historical tournament analysis dataset.','Third-party datasets need separate freshness, schema and license review.')
    ],[1.45,2.8,2.55])
    d.add_page_break()
    d.add_heading('2. Dataset design',1)
    d.add_paragraph('The target archive is a traceable data lake of raw-source snapshots, normalized tables, assets, and metadata. Each extracted row should retain the source URL, retrieval timestamp, coverage state and confidence where possible.')
    add_table(d,['Layer','Contents','Example files'],[
        ('Raw source','Original HTML and downloaded source archives preserved for audit.','raw_html/, reference/, raw_sources/'),
        ('Parsed tables','HTML tables saved individually as UTF-8 CSV, plus discovered links.','parsed_tables/*_table_*.csv; *_links.csv'),
        ('Hero analytics','Picks, bans, wins/losses, rates, tournament, patch and games with recorded draft.','templates/hero_stats_schema.csv'),
        ('Game-level drafts','Series/game ID, teams, winner, duration, patch, 10 picks, ban strings and pick order.','templates/game_level_schema.csv'),
        ('Coverage and provenance','URL status, file bytes, SHA-256 digest, timestamp, row count and extraction limitations.','source_fetch_log.csv; extraction_report.json')
    ],[1.25,3.15,2.4])
    d.add_heading('Counter analysis that can be derived',2)
    for x in [
        'Hero vs hero outcomes by patch, tournament tier and region, with sample size and confidence interval.',
        'Pick/ban priority and presence, separated from win rate to avoid confusing draft pressure with effectiveness.',
        'Team composition patterns, frequent pairings, ban responses and counter-pick outcomes when draft order is captured.',
        'Season comparison for each MPL region and event; stale patches should not be mixed with current meta without a patch control.'
    ]:
        d.add_paragraph(x,style='List Bullet')
    d.add_page_break()
    d.add_heading('3. Existing sourced snapshots',1)
    d.add_paragraph('These CSV files are included as starter data. They are source-attributed snapshots, not a completed full extraction.')
    add_table(d,['CSV','Scope','Records in current file','How to interpret'],[
        ('hero_meta_global_snapshot.csv','Global pro 20-championship sample','25 hero rows','Pick count, ban count, and reported pick-win-rate percentage.'),
        ('msc_2026_hero_snapshot.csv','MSC 2026','12 hero rows','Event hero snapshot; includes picks, bans, wins/losses where captured.'),
        ('mpl_my_s18_visible_snapshot.csv','MPL Malaysia S18','7 hero rows','Visible subset only; do not infer that omitted heroes had zero picks.'),
        ('tournament_reference_index.csv','Current visible event list','5 event rows','Displayed game counts from index snapshot.')
    ],[2.05,1.55,1.2,2.0])
    d.add_heading('Sample observations — not a universal tier list',2)
    d.add_paragraph('In the captured MSC 2026 sample, Atlas showed 44 picks and 65 bans with a 65.9% win rate across its 29 wins and 15 losses; Freya showed 19 picks, 90 bans and a 73.7% win rate across 14 wins and 5 losses. Their low pick sample sizes and ban pressure matter: these values alone do not establish a universal best hero or a direct counter relation.')
    d.add_paragraph('For the visible global pro sample, Freya has 1,993 bans and 432 picks, while Zhuxin has 809 picks and 1,233 bans. These figures indicate draft presence in that sample, not necessarily a pick-rate denominator that is comparable across all tournaments.')
    d.add_heading('4. How to execute the complete collection',1)
    for x in [
        'Upload the package contents, including the hidden .github/workflows folder, to a GitHub repository; open Actions and run “MLBB Full Public Data Extraction”. This can be done from a phone browser, subject to GitHub account limits.',
        'Keep asset downloads enabled to capture open reference JSON and visual assets. The collector follows public links discovered from the index and respects a delay between requests.',
        'Download the artifact and inspect extraction_report.json, source_fetch_log.csv, tournament_page_inventory.csv and source_archives_manifest.json.',
        'Audit every region/season. Mark a season complete only when required game counts, draft coverage, source validity and duplicate checks have passed.',
        'If output size is under 100 MiB, check accessible sources and logs. Do not manufacture rows or add arbitrary filler files to meet the threshold.'
    ]: d.add_paragraph(x,style='List Number')
    d.add_page_break()
    d.add_heading('5. Quality gates and limitations',1)
    add_table(d,['Gate','Must pass before calling the dataset complete'],[
        ('Coverage','Every intended MPL region/season, M-Series, MSC and other event is present with counts matched to a cited source.'),
        ('Game integrity','Unique stable game IDs or robust composite keys; no duplicate game rows; winners match sourced results.'),
        ('Draft integrity','10 picks per completed game, bans/order retained where source provides them; unknown values stay null.'),
        ('Rate calculations','Wins / picks and event-game denominators clearly specified; pick win rate not confused with all-game win rate.'),
        ('Patch/era control','Patch, season and event dates tracked; comparisons across patch changes are labelled.'),
        ('Provenance','Every data table has a source URL, retrieval timestamp and collection log. License requirements are reviewed.'),
        ('Size threshold','Actual measured extracted output exceeds 100 MiB; no padding, duplicate copies or synthetic records.'),
    ],[1.4,5.6])
    d.add_heading('6. References and license notes',1)
    refs=[
        'MLBBDex esports index — https://mlbbdex.com/en/esports (accessed 22 October 2026).',
        'MLBBDex MSC 2026 — https://mlbbdex.com/en/esports/msc-2026.',
        'MLBB Analytics — https://www.mlbbanalytics.com/heroes and https://www.mlbbanalytics.com/esports.',
        'MLBBHub MPL MY hero stats — https://mlbbhub.com/mpl/my/stats.',
        'Liquipedia Statistics Portal — https://liquipedia.net/mobilelegends/Portal:Statistics.',
        'LiTianYeoh tournament analysis — https://github.com/LiTianYeoh/MLBB_Tournament_Analysis (README documents consolidated file dated 2024-05-05).',
        'Ceplin03 MLBB data and assets — https://github.com/Ceplin03/database-mlbb.Mobile-Legends-Bang-Bang (repository README states MIT; inspect each license/asset before redistribution).'
    ]
    for ref in refs: d.add_paragraph(ref,style='List Bullet')
    p=d.add_paragraph(); p.paragraph_format.space_before=Pt(8)
    r=p.add_run('Final status: V3 collection workflow prepared; exhaustive network collection and 100 MiB output are pending a successful run in an internet-enabled environment.'); r.bold=True; r.font.color.rgb=RGBColor.from_string('B45025')
    d.save(REPORT)


def add_bg(slide, color=NAVY):
    fill=slide.background.fill; fill.solid(); fill.fore_color.rgb=PPTColor.from_string(color)
def add_text(slide,x,y,w,h,text,size=18,color=WHITE,bold=False,font='Aptos',align=None):
    box=slide.shapes.add_textbox(PInches(x),PInches(y),PInches(w),PInches(h)); tf=box.text_frame; tf.clear(); tf.word_wrap=True
    for i,line in enumerate(str(text).split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line; p.font.name=font; p.font.size=PPt(size); p.font.bold=bold; p.font.color.rgb=PPTColor.from_string(color)
        if align is not None: p.alignment=align
    return box
def add_accent(slide,x,y,w,h,color=CYAN):
    sh=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,PInches(x),PInches(y),PInches(w),PInches(h)); sh.fill.solid(); sh.fill.fore_color.rgb=PPTColor.from_string(color); sh.line.fill.background(); return sh
def add_bullet_slide(prs,title,subtitle,bullets,idx):
    s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s, NAVY if idx%2 else '142D43'); add_accent(s,.65,.55,.12,.55)
    add_text(s,.95,.5,11.7,.55,title,26,WHITE,True,'Aptos Display')
    add_text(s,.95,1.18,11.3,.45,subtitle,11,CYAN,False)
    y=1.85
    for i,b in enumerate(bullets):
        add_accent(s,.98,y+.06,.11,.11,ORANGE if i==0 else CYAN)
        add_text(s,1.25,y,10.8,.73,b,16,WHITE,False)
        y+=.93
    add_text(s,.7,7.08,11.8,.22,'MLBB META & COUNTER INTELLIGENCE  •  V3  •  22 OCT 2026',7,'9EB0BF',False,align=PP_ALIGN.RIGHT)
    return s

def make_presentation():
    p=Presentation(); p.slide_width=PInches(13.333); p.slide_height=PInches(7.5)
    s=p.slides.add_slide(p.slide_layouts[6]); add_bg(s); add_accent(s,.8,.8,.14,1.1)
    add_text(s,1.15,.78,10.9,.8,'MLBB META & COUNTER',34,WHITE,True,'Aptos Display')
    add_text(s,1.15,1.68,10.7,.6,'FULL PUBLIC-DATA EXTRACTION • V3',22,CYAN,True)
    add_text(s,1.15,2.65,9.8,1.1,'A reproducible route to tournament, hero, pick/ban and counter analysis across MPL, MSC and M-Series.',19,'D7E5EF')
    add_accent(s,1.15,4.15,4.0,.05,ORANGE)
    add_text(s,1.15,4.45,10.8,1.1,'Integrity note: this package contains a collector plus sourced starter snapshots. Exhaustive live extraction and the 100 MiB target have not been verified in this offline build.',14,'F3C9A2')
    add_text(s,.8,7.08,11.8,.22,'RESEARCH BUILD  •  22 OCTOBER 2026',7,'9EB0BF',False,align=PP_ALIGN.RIGHT)
    add_bullet_slide(p,'What V3 includes','Collector framework plus currently captured starter data',[
        'Sourced CSV snapshots: global pro hero stats, MSC 2026, visible MPL Malaysia S18 rows and tournament index counts.',
        'Extractor saves source HTML, parses public tables, discovers linked event pages and fetches open reference repositories/assets.',
        'Provenance output: source URL, retrieval time, HTTP status, byte count, SHA-256, page inventory and run-level report.',
        'Word report, this deck, CSV schemas, mobile run instructions and manual GitHub Actions workflow.'
    ],1)
    add_bullet_slide(p,'Visible tournament index snapshot','Counts shown by MLBBDex on 22 October 2026',[
        'MPL Philippines S18: 134 games listed; MPL Indonesia S18: 174.',
        'MPL Malaysia S18: 136; MSC 2026: 110; M7 World Championship: 137.',
        'The site reports 554 games from MSC 2026 and the three live S18 MPL leagues within its recent 120-day view.',
        'Index counts do not prove full game-draft completeness or exhaustive historical coverage.'
    ],2)
    add_bullet_slide(p,'Data model','Separate hero aggregates from match-level records',[
        'Hero aggregates: pick, ban, wins/losses, rate denominators, event, region, season and patch.',
        'Game rows: stable game/series identity, teams, date, patch, winner, duration, picks, bans and pick order where available.',
        'Keep unknown values null; retain source URL and collection timestamp.',
        'Do not infer “zero picks” from a hero missing in a partial table.'
    ],3)
    add_bullet_slide(p,'Counter-analysis outputs','Evidence-based rather than static tier labels',[
        'Hero-vs-hero outcome matrices grouped by patch and event tier, with sample sizes.',
        'Pick/ban priority and presence separated from win rate.',
        'Draft pairings, role/lane combinations and first-pick/red-side effects when the source retains order.',
        'Patch controls prevent stale drafts from being mixed into current-meta recommendations.'
    ],4)
    add_bullet_slide(p,'Sourced snapshot examples','Illustrative only; not a universal tier list',[
        'MSC 2026 snapshot: Atlas 44 picks, 65 bans, 65.9% WR (29–15); Freya 19 picks, 90 bans, 73.7% WR (14–5).',
        'Global pro snapshot: Freya 432 picks / 1,993 bans; Zhuxin 809 picks / 1,233 bans in the captured source sample.',
        'Small pick samples can swing win rate; high bans show draft pressure, not proof that a hero hard-counters another.',
        'The CSVs include source URLs and snapshot dates for audit.'
    ],5)
    add_bullet_slide(p,'How to run from a phone','GitHub Actions workflow is included',[
        'Upload the package contents to a GitHub repository, including .github/workflows/extract.yml.',
        'Open Actions → “MLBB Full Public Data Extraction” → Run workflow.',
        'Set a page limit, keep reference assets enabled, and require a 100 MiB measured output.',
        'Download the artifact and inspect extraction_report.json and source_fetch_log.csv; pages that fail are recorded, not bypassed.'
    ],6)
    add_bullet_slide(p,'100 MiB policy and limits','No padding, duplicates or synthetic match rows',[
        '100 MiB is checked against real downloaded/parsed data and legitimate reference assets.',
        'If online sources yield less than 100 MiB, the run reports that target was not met instead of manufacturing size.',
        'Public source access, rate limits, site changes, historical gaps and license terms can restrict collection.',
        'Only mark a season complete after counts, duplicate checks, draft coverage and provenance pass the quality gates.'
    ],7)
    add_bullet_slide(p,'Primary sources','Review licenses and source coverage before redistribution',[
        'MLBBDex: mlbbdex.com/en/esports (its index says tournament data is adapted from Liquipedia under CC BY-SA 3.0).',
        'MLBB Analytics: mlbbanalytics.com/heroes and /esports; MLBBHub: mlbbhub.com/mpl/my/stats.',
        'Historical scrape project: github.com/LiTianYeoh/MLBB_Tournament_Analysis (its README references a 2024-05-05 consolidated dataset).',
        'Community reference data/assets: github.com/Ceplin03/database-mlbb.Mobile-Legends-Bang-Bang (README states MIT; inspect license files).'
    ],8)
    p.save(PRES)

if __name__=='__main__':
    make_report(); make_presentation(); print(REPORT); print(PRES)
